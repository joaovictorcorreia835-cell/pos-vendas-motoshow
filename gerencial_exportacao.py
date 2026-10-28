"""Excel e PDF gerenciais utilizam a mesma leitura filtrada da tela."""
from io import BytesIO
from datetime import date, datetime
from flask import send_file
from gerencial_tabelas import tabelas, METRICAS

EXTRAS = [(key,label,tipo,'','') for key,label,tipo in [
    ('bruto','Bruto vendido no balcão','money'),('liquido_vendas','Líquido das vendas antes das devoluções','money'),
    ('descontos','Descontos concedidos','money'),('desconto_pct','Desconto médio ponderado','percent'),
    ('maior_desconto','Maior desconto','percent'),('vendas_desconto','Vendas com desconto','int'),
    ('devolucoes','Quantidade de devoluções','int'),('valor_devolvido','Valor devolvido','money'),
    ('impacto_margem','Impacto das devoluções na margem','money'),('devolucao_pct','Devoluções sobre vendas','percent')]]


def secoes(d):
    matriz = [dict(modelo=r['modelo'], **{key:len(value) for key,value in r['celulas'].items()}) for r in d['matriz']]
    return tabelas(d) + [
        dict(id='aging',titulo='Aging da carteira em aberto',rows=d['aging'],cols=[
            ('label','Faixa','text'),('quantidade','OS','int'),('pecas','Peças','money'),('servicos','Serviços','money'),('total','Carteira em aberto','money')]),
        dict(id='quadrantes',titulo='Faturamento e margem por consultor',rows=d['quadrantes'],cols=[
            ('consultor','Consultor','text'),('balcao','Balcão líquido','money'),('margem','Margem','money'),('margem_pct','Margem %','percent'),('quadrante','Situação relativa ao grupo','text')]),
        dict(id='matriz',titulo='Modelos e tipos de serviço',rows=matriz,cols=[('modelo','Modelo','text')]+[(k,k,'int') for k in ['Revisão','Garantia','Serviços','Seguro','Outros']])]


def metadados(d):
    return [
        ('Período', d['inicio'].isoformat() + ' a ' + d['fim'].isoformat()),
        ('Emissão', d['emitido_em']),
        ('Unidade', d['filtros'].get('unidade') or 'Todas'),
        ('Consultor', d['filtros'].get('consultor') or 'Todos'),
        ('Tipo de OS', d['filtros'].get('tipo') or 'Todos'),
        ('Modelo', d['filtros'].get('modelo') or 'Todos'),
        ('Origem', 'Somente transações importadas; sem faturamento manual.'),
        ('Datas fiscais', 'OS sem data fiscal diária só entram no financeiro de meses completos.'),
        ('Margem', 'Faturado menos custo informado; não é lucro líquido após despesas.'),
        ('Carteira', 'Situação atual das OS com entrada no período. Idade calculada até hoje.'),
        ('Clientes', 'CPF/CNPJ quando disponível; caso contrário, ID de cadastro. Sem união por nome.'),
        ('Limites', 'Metas e itens individuais de peças indisponíveis; comparações dependem de histórico.'),
        ('Registros processados', str(d['auditoria']['processados'])),
        ('Registros já existentes nas fontes',str(d['auditoria']['duplicados'])),
        ('Auditoria de erros',d['auditoria']['erros'])]


def excel(d):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = 'Dados e filtros'
    for key, value in metadados(d):
        ws.append([key, value])
    ws.append(['Indicador', 'Valor'])
    for key, label, tipo, _, _ in METRICAS+EXTRAS:
        value = d['indicadores'][key]
        ws.append([label, value / 100 if tipo == 'money' and value is not None else value])
        if tipo == 'money':
            ws.cell(ws.max_row, 2).number_format = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
    for tabela in secoes(d) + [dict(id='OS importadas',cols=[
        ('numero','OS','text'),('cliente','Cliente','text'),('modelo','Modelo','text'),
        ('tipo','Tipo','text'),('consultor','Consultor','text'),('entrada','Entrada','date'),
        ('status','Status','text'),('competencia','Competência fiscal','text'),
        ('pecas','Peças','money'),('servicos','Serviços','money'),('total','Valor OS','money')],rows=d['registros'])]:
        ws = wb.create_sheet(tabela['id'][:31])
        ws.append([label for _, label, _ in tabela['cols']])
        for row in tabela['rows']:
            values = []
            for key, _, tipo in tabela['cols']:
                value = row.get(key)
                values.append(value / 100 if tipo == 'money' and value is not None else value)
            ws.append(values)
            for i, (_, _, tipo) in enumerate(tabela['cols'], 1):
                if tipo == 'money':
                    ws.cell(ws.max_row, i).number_format = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
                elif tipo == 'date':
                    ws.cell(ws.max_row, i).number_format = 'dd/mm/yyyy'
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    ws = wb.create_sheet('Fontes')
    ws.append(['Relatório','SHA-256','Última leitura'])
    for row in d['auditoria']['fontes']:
        ws.append([row['arquivo'],row['sha256'],row['ultima']])
    for ws in wb:
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='082B40')
        for row in ws:
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type='s'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for i in range(1,ws.max_column+1):
            ws.column_dimensions[get_column_letter(i)].width=min(48,max(16,max(len(str(c.value or '')) for c in ws[get_column_letter(i)])+2))
    stream=BytesIO()
    wb.save(stream)
    stream.seek(0)
    return send_file(stream,as_attachment=True,download_name='painel_gerencial.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def formato(value, tipo):
    if value is None:
        return 'Não disponível'
    if tipo=='money':
        return 'R$ '+f'{value/100:,.2f}'.replace(',','X').replace('.',',').replace('X','.')
    if tipo=='percent':
        return f'{value:.2f}%'.replace('.',',')
    if isinstance(value,(date,datetime)):
        return value.strftime('%d/%m/%Y')
    return str(value)


def pdf(d):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from xml.sax.saxutils import escape
    stream=BytesIO()
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='CelulaGerencial',fontSize=6.5,leading=8,textColor=colors.HexColor('#102e43')))
    story=[Paragraph('PAINEL GERENCIAL - MOTOSHOW',styles['Title'])]
    for key,value in metadados(d):
        story.append(Paragraph(escape(key+': '+str(value)),styles['BodyText']))
    def tabela(headers,rows):
        values=[[Paragraph(escape(str(v)),styles['CelulaGerencial']) for v in row] for row in [headers]+rows]
        t=Table(values,repeatRows=1,colWidths=[770/len(headers)]*len(headers),hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dbe9f2')),
            ('GRID',(0,0),(-1,-1),.25,colors.HexColor('#9db2c2')),('VALIGN',(0,0),(-1,-1),'TOP'),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f3f7fa')])]))
        story.append(t)
    story.append(Spacer(1,14))
    tabela(['Indicador','Valor'],[[label,formato(d['indicadores'][key],tipo)] for key,label,tipo,_,_ in METRICAS+EXTRAS])
    for section in secoes(d):
        story.extend([PageBreak(),Paragraph(escape(section['titulo']),styles['Heading1'])])
        if section['rows']:
            tabela([label for _,label,_ in section['cols']],[[formato(row.get(key),tipo) for key,_,tipo in section['cols']] for row in section['rows']])
        else:
            story.append(Paragraph('Nenhum registro para os filtros selecionados.',styles['BodyText']))
    story.extend([PageBreak(),Paragraph('Dados e fontes do painel',styles['Heading1'])])
    tabela(['Relatório','SHA-256','Leitura'],[[r['arquivo'],r['sha256'],r['ultima']] for r in d['auditoria']['fontes']])
    for divergencia in d['auditoria']['divergencias']:
        story.append(Spacer(1,10))
        story.append(Paragraph(escape(divergencia['arquivo']+': diferença entre margem do rodapé e soma das linhas: '+formato(divergencia['diferenca'],'money')+'. Foram usadas as transações detalhadas.'),styles['BodyText']))
    def rodape(canvas,doc):
        canvas.setFont('Helvetica',8)
        canvas.drawString(28,18,'Motoshow - '+d['inicio'].isoformat()+' a '+d['fim'].isoformat())
        canvas.drawRightString(812,18,'Pagina '+str(doc.page))
    SimpleDocTemplate(stream,pagesize=landscape(A4),leftMargin=28,rightMargin=28,topMargin=28,bottomMargin=32).build(story,onFirstPage=rodape,onLaterPages=rodape)
    stream.seek(0)
    return send_file(stream,as_attachment=True,download_name='painel_gerencial.pdf',mimetype='application/pdf')
