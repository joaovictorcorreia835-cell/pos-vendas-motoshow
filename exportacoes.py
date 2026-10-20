"""Exportações Excel das listas operacionais, preservando filtros e tipos de dados."""
from datetime import datetime
from io import BytesIO
from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from indicadores import prazo,hoje_local

MONEY='"R$" #,##0.00;[Red]-"R$" #,##0.00'
def workbook_response(title,headers,rows,filename,money_columns,date_columns,total_columns):
 wb=Workbook();ws=wb.active;ws.title=title
 ws.append(headers)
 for values in rows:
  ws.append(values)
  # Dados importados são texto literal, inclusive quando começam por =, +, - ou @.
  for cell in ws[ws.max_row]:
   if isinstance(cell.value,str):cell.data_type='s'
 last=ws.max_row
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='082B40');cell.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=30
 for row in ws.iter_rows(min_row=2):
  for cell in row:
   if cell.column in money_columns:cell.number_format=MONEY
   elif cell.column in date_columns:cell.number_format='dd/mm/yyyy hh:mm' if isinstance(cell.value,datetime) else 'dd/mm/yyyy'
   cell.alignment=Alignment(vertical='top',wrap_text=True)
 if rows:
  totalrow=last+1;ws.cell(totalrow,1,'TOTAL')
  for col in total_columns:ws.cell(totalrow,col,f'=SUM({get_column_letter(col)}2:{get_column_letter(col)}{last})').number_format=MONEY
  for cell in ws[totalrow]:cell.font=Font(bold=True,color='082B40');cell.fill=PatternFill('solid',fgColor='DEEBF7')
 ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}{last}'
 for col,header in enumerate(headers,1):
  width=18 if col in money_columns or col in date_columns else min(45,max(14,len(header)+3,max((len(str(r[col-1] or '')) for r in rows),default=0)+2))
  ws.column_dimensions[get_column_letter(col)].width=width
 ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:1'
 ws.oddFooter.center.text='Página &P de &N'
 stream=BytesIO();wb.save(stream);stream.seek(0)
 return send_file(stream,as_attachment=True,download_name=filename,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

def export_pedidos(rows,mode=''):
 data=[]
 for p in rows:
  age=prazo(p.inserido_em,p.status=='Aberto')
  data.append([p.numero,p.codigo_interno,p.inserido_em,p.cliente,p.dados.get('Telefone',''),p.consultor,p.status,age['rotulo'],age['dias'],p.validade,p.bruto_centavos/100,p.liquido_centavos/100,p.condicao,p.observacao,p.unidade])
 name='pedidos_atrasados' if mode=='atrasado' else 'pedidos_proximos_vencimento' if mode=='atencao' else 'pedidos'
 return workbook_response('Pedidos',['Pedido','Código interno','Inserção','Cliente','Telefone','Consultor','Situação','Prazo','Dias em aberto','Validade','Valor bruto','Valor líquido','Condição de pagamento','Observação','Unidade'],data,f'{name}_{hoje_local().isoformat()}.xlsx',{11,12},{3,10},{11,12})

def export_ordens(rows,mode=''):
 data=[]
 for p in rows:
  age=prazo(p.data,p.status not in ['Finalizada','Cancelada'])
  labor=sum(i.valor_liquido for i in p.itens if i.categoria=='Mão de obra')
  data.append([p.ordem.numero,p.data,p.cliente.nome,p.cliente.telefone,p.modelo,p.placa,p.consultor.nome if p.consultor else '',p.mecanico.nome if p.mecanico else '',p.tipo,p.status,age['rotulo'],age['dias'],(p.total-labor)/100,labor/100,p.total/100,p.unidade])
 name='os_atrasadas' if mode=='atrasado' else 'os_proximas_vencimento' if mode=='atencao' else 'ordens_servico'
 return workbook_response('Ordens de serviço',['OS','Entrada','Cliente','Telefone','Moto','Placa','Consultor','Mecânico','Tipo','Status','Prazo','Dias em aberto','Peças e outros','Mão de obra','Valor total','Unidade'],data,f'{name}_{hoje_local().isoformat()}.xlsx',{13,14,15},{2},{13,14,15})
