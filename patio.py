"""Consulta e atualização do estoque de motos em pátio."""
import csv
import io
import hashlib
import json
import unicodedata
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

from flask import Blueprint, abort, flash, redirect, render_template, request, url_for
from openpyxl import load_workbook
from models import db

bp = Blueprint('patio', __name__)


class MotoPatio(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    chassi = db.Column(db.String(80), unique=True, nullable=False)
    modelo = db.Column(db.String(200), nullable=False)
    placa = db.Column(db.String(30), default='')
    cor = db.Column(db.String(100), default='')
    unidade = db.Column(db.String(150), default='')
    condicao = db.Column(db.String(100), default='')
    situacao = db.Column(db.String(150), nullable=False)
    faturado_para = db.Column(db.String(200), default='')
    atualizado_em = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)
    arquivo = db.Column(db.String(255), nullable=False)


class RelatorioPatio(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    fonte = db.Column(db.String(30), nullable=False)
    unidade = db.Column(db.String(150), nullable=False)
    sha256 = db.Column(db.String(64), nullable=False, unique=True)
    arquivo = db.Column(db.String(255), nullable=False)
    emitido_em = db.Column(db.DateTime, nullable=False)
    importado_em = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)
    dados = db.Column(db.Text, nullable=False)


class LocalizacaoMoto(db.Model):
    chassi = db.Column(db.String(80), primary_key=True)
    localizacao = db.Column(db.String(30), nullable=False)
    atualizado_em = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)


def importar_pdfs(uploads):
    """Valida todos os arquivos antes de alterar o banco; o chamador confirma."""
    from patio_pdf import parse_pdf
    prepared = []
    for upload in uploads:
        raw = upload.read()
        data = parse_pdf(raw)
        prepared.append((upload.filename, hashlib.sha256(raw).hexdigest(), data))
    count = 0
    for filename, digest, data in prepared:
        if db.session.scalar(db.select(RelatorioPatio.id).where(RelatorioPatio.sha256 == digest)):
            continue
        latest = db.session.scalar(db.select(RelatorioPatio).where(RelatorioPatio.fonte == data['fonte'], RelatorioPatio.unidade == data['unidade']).order_by(RelatorioPatio.emitido_em.desc(), RelatorioPatio.id.desc()))
        emitted = datetime.fromisoformat(data['emitido_em'])
        if latest and latest.emitido_em > emitted:
            raise ValueError('O relatório enviado é mais antigo que o já importado para esta unidade.')
        db.session.add(RelatorioPatio(fonte=data['fonte'], unidade=data['unidade'], sha256=digest,
            arquivo=Path(filename.replace('\\', '/')).name[:255], emitido_em=emitted,
            dados=json.dumps(data, ensure_ascii=False)))
        count += 1
    return count


def estoque_atual():
    snapshots = {}
    for source in db.session.scalars(db.select(RelatorioPatio).order_by(RelatorioPatio.emitido_em.desc(), RelatorioPatio.id.desc())):
        snapshots.setdefault((source.fonte, source.unidade), source)
    entries = {}
    for moto in db.session.scalars(db.select(MotoPatio)):
        row = {c.name: getattr(moto, c.name) for c in MotoPatio.__table__.columns}
        row.update(presenca='Informada na planilha', parcial=False, ano='', dias=None, custo_centavos=None, compra_centavos=None, oferta_centavos=None, fornecedor='', titulo='', secao='', codigo='', negociacao='', emitido_em=None)
        entries[moto.chassi] = row
    reports = []
    for (kind, unit), source in snapshots.items():
        data = json.loads(source.dados)
        reports.append(dict(arquivo=source.arquivo, fonte=kind, emitido_em=source.emitido_em, quantidade=data['quantidade']))
        if kind != 'oferta':
            continue
        for row in data['rows']:
            row.update(arquivo=source.arquivo, atualizado_em=source.importado_em, emitido_em=source.emitido_em, parcial=False, titulo='', secao='Posição de veículos em oferta')
            entries[row['chassi']] = row
    pending = []
    transito = []
    for (kind, unit), source in snapshots.items():
        if kind != 'posicao':
            continue
        for record in json.loads(source.dados)['rows']:
            if record['secao'] == 'MOTOS EM TRANSITO A PAGAR':
                transito.append(dict(record, unidade=unit, arquivo=source.arquivo, emitido_em=source.emitido_em))
            if 'MOTOSHOW' not in record['titular'].upper() or 'FATURADAS' not in record['secao']:
                continue
            # Os chassis cortados no PDF financeiro não identificam uma moto.
            # Nunca vincular por prefixo nem inventar o dígito ausente.
            full_match = entries.get(record['chassi']) if len(record['chassi']) == 17 else None
            if full_match:
                full_match['faturado_para'] = record['titular']
                full_match['titulo'] = record['titulo']
                continue
            pending.append(dict(chassi=record['chassi'], modelo='Modelo não informado', placa='', cor='',
                unidade=unit, condicao='Nova', situacao='Faturada em nome da Motoshow',
                faturado_para=record['titular'], presenca='Pátio a confirmar', parcial=len(record['chassi']) != 17,
                arquivo=source.arquivo, atualizado_em=source.importado_em, emitido_em=source.emitido_em,
                ano='', dias=None, compra_centavos=None, custo_centavos=None, oferta_centavos=None,
                fornecedor='', titulo=record['titulo'], secao=record['secao'], codigo='', negociacao=''))
    result = [SimpleNamespace(**row) for row in list(entries.values()) + pending]
    locations = {m.chassi: m.localizacao for m in db.session.scalars(db.select(LocalizacaoMoto))}
    for moto in result:
        moto.localizacao = locations.get(moto.chassi, 'A confirmar') if not moto.parcial else 'A confirmar'
        moto.possivel_transito = any(moto.unidade == t['unidade'] and moto.chassi.startswith(t['chassi']) for t in transito) if not moto.parcial else False
    # O mestre é persistente: ausência em relatório não apaga uma moto.
    from inventario_models import MotoEstoque
    from inventario_service import normalizar_chassi
    from sqlalchemy.orm import selectinload
    normalized={}; others=[]
    for dto in result:
        if dto.parcial:
            others.append(dto); continue
        try:
            dto.chassi=normalizar_chassi(dto.chassi)
            normalized.setdefault(dto.chassi,dto)
        except ValueError:
            dto.parcial=True
            dto.localizacao='A confirmar'
            others.append(dto)
    result=list(normalized.values())+others
    masters=db.session.scalars(db.select(MotoEstoque).options(selectinload(MotoEstoque.localizacao))).all()
    by_vin={m.chassi:m for m in result if not m.parcial}
    for master in masters:
        if not master.ativo or master.saida_fisica_confirmada:
            result=[m for m in result if m.parcial or m.chassi!=master.chassi_normalizado]
            continue
        dto=by_vin.get(master.chassi_normalizado)
        if dto is None:
            dto=SimpleNamespace(chassi=master.chassi_normalizado,modelo=master.modelo,placa=master.placa,cor=master.cor,
                unidade=master.unidade_nome,condicao='Nova',situacao=master.status_comercial,faturado_para='',presenca='Cadastro mestre',
                parcial=False,arquivo=master.origem_dado,atualizado_em=master.atualizado_em,emitido_em=None,ano=master.ano_modelo,
                dias=master.dias_estoque,compra_centavos=None,custo_centavos=None,oferta_centavos=None,fornecedor='',titulo='',secao='',codigo='',negociacao='',possivel_transito=False)
            result.append(dto)
        dto.mestre_id=master.id; dto.modelo=master.modelo; dto.cor=master.cor; dto.unidade=master.unidade_nome
        dto.dias=master.dias_estoque
        dto.localizacao='Dentro da oficina' if master.localizacao.grupo=='OFICINA' else 'Em trânsito' if master.localizacao.grupo=='TRANSITO' else 'A confirmar' if master.localizacao.grupo=='A_CONFIRMAR' else master.localizacao.nome
        dto.situacao={'DISPONÍVEL':'Disponível para venda','AGUARDANDO APROVAÇÃO':'Aguardando aprovação'}.get(master.status_comercial,master.status_comercial)
    for report in reports:
        if report['fonte'] == 'posicao':
            report['transito'] = [t for t in transito if t['arquivo'] == report['arquivo']]
    return sorted(result, key=lambda m: (m.modelo, m.chassi)), reports


from inventario import mutacao


@bp.post('/patio/localizacao')
@mutacao('OPERADOR')
def atualizar_localizacao():
    from inventario_service import permitir
    permitir('OPERADOR')
    from inventario_service import normalizar_chassi
    try:
        chassi=normalizar_chassi(request.form.get('chassi',''))
    except ValueError:
        abort(400,description='A localização exige chassi completo válido.')
    location = request.form.get('localizacao', '')
    if location not in ('Dentro da oficina', 'Em trânsito', 'A confirmar'):
        abort(400)
    rows, _ = estoque_atual()
    if not any(m.chassi == chassi and not m.parcial for m in rows):
        abort(400, description='Confirme a localização usando um chassi completo do estoque.')
    row = db.session.get(LocalizacaoMoto, chassi)
    if row is None:
        row = LocalizacaoMoto(chassi=chassi)
        db.session.add(row)
    row.localizacao = location
    row.atualizado_em = datetime.utcnow()
    from inventario_models import MotoEstoque
    from inventario_service import mover, obter_moto
    master=db.session.scalar(db.select(MotoEstoque).where(MotoEstoque.chassi_normalizado==chassi))
    if not master:
        abort(400,description='Cadastre o chassi completo no estoque mestre antes de confirmar sua presença.')
    mover(obter_moto(master.id),{'Dentro da oficina':'OFICINA','Em trânsito':'TRANSITO','A confirmar':'A_CONFIRMAR'}[location], 'Confirmação pela tela legada')
    db.session.commit()
    flash('Localização da moto atualizada.', 'success')
    return redirect(url_for('patio.listagem', **{k: request.form[k] for k in ('q', 'grupo', 'unidade', 'situacao', 'localizacao') if request.form.get(k) and k != 'localizacao'}))


def normalizar(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(value or '').strip().lower()) if not unicodedata.combining(c))


def atrasada(moto):
    return moto.dias is not None and moto.dias > 45 and getattr(moto, 'localizacao', '') != 'Em trânsito'


def resumo_atrasadas(unidade=''):
    rows, reports = estoque_atual()
    quantidade = sum(atrasada(m) for m in rows if not unidade or m.unidade == unidade)
    query = dict(grupo='atrasadas')
    if unidade:
        query['unidade'] = unidade
    datas = [r['emitido_em'] for r in reports if r['fonte'] == 'oferta']
    nota = ('Entrada do mestre baseada no relatório de ' + max(datas).strftime('%d/%m/%Y') + '; dias corridos de estoque. Presença física depende de conferência'
            if datas else 'Importe o relatório de veículos em oferta para atualizar a quantidade.')
    return dict(quantidade=quantidade, url=url_for('patio.listagem', **query), nota=nota)


def ler_relatorio(upload):
    extension = Path(upload.filename or '').suffix.lower()
    if extension == '.xlsx':
        book = load_workbook(upload.stream, read_only=True, data_only=True)
        try:
            rows = list(book.active.iter_rows(values_only=True))
        finally:
            book.close()
    elif extension == '.csv':
        raw = upload.read()
        try:
            content = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            content = raw.decode('cp1252')
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=';,\t')
        except csv.Error:
            raise ValueError('Não foi possível identificar as colunas do CSV.')
        rows = list(csv.reader(io.StringIO(content), dialect))
    else:
        raise ValueError('Selecione um relatório Excel (.xlsx) ou CSV.')
    rows = [row for row in rows if any(str(v or '').strip() for v in row)]
    if len(rows) < 2:
        raise ValueError('O relatório deve conter cabeçalhos e pelo menos uma moto.')
    aliases = {'chassi': 'chassi', 'chassis': 'chassi', 'modelo': 'modelo', 'placa': 'placa', 'cor': 'cor', 'unidade': 'unidade', 'condicao': 'condicao', 'situacao': 'situacao', 'status': 'situacao', 'faturado para': 'faturado_para', 'faturado em nome de': 'faturado_para'}
    headers = [aliases.get(normalizar(v)) for v in rows[0]]
    if not {'chassi', 'modelo', 'situacao'}.issubset(headers):
        raise ValueError('O relatório precisa das colunas Chassi, Modelo e Situação. Colunas opcionais: Placa, Cor, Unidade, Condição e Faturado para.')
    if len([h for h in headers if h]) != len(set(h for h in headers if h)):
        raise ValueError('O relatório contém colunas repetidas.')
    result, seen = [], set()
    for number, row in enumerate(rows[1:], 2):
        data = {key: str(row[i] if i < len(row) and row[i] is not None else '').strip() for i, key in enumerate(headers) if key}
        data['chassi'] = data['chassi'].upper()
        if any(not data[key] for key in ('chassi', 'modelo', 'situacao')):
            raise ValueError(f'Linha {number}: preencha Chassi, Modelo e Situação.')
        from inventario_service import normalizar_chassi
        try:
            data['chassi']=normalizar_chassi(data['chassi'])
        except ValueError as error:
            raise ValueError(f'Linha {number}: {error}') from None
        if data['chassi'] in seen:
            raise ValueError(f'Linha {number}: chassi repetido no relatório.')
        for key, limit in [('chassi',80),('modelo',200),('placa',30),('cor',100),('unidade',150),('condicao',100),('situacao',150),('faturado_para',200)]:
            if len(data.get(key, '')) > limit:
                raise ValueError(f'Linha {number}: o campo {key} é muito longo.')
            data.setdefault(key, '')
        seen.add(data['chassi'])
        result.append(data)
    return result


@bp.route('/patio', methods=['GET', 'POST'])
def listagem():
    if request.method == 'POST':
        from inventario_service import permitir
        permitir('GESTOR')
        uploads = [u for u in request.files.getlist('arquivo') if u.filename]
        try:
            if not uploads:
                raise ValueError('Selecione o relatório para importar.')
            if all(Path(u.filename).suffix.lower() == '.pdf' for u in uploads):
                count = importar_pdfs(uploads)
                from inventario_service import sincronizar_fontes
                db.session.flush(); sincronizar_fontes()
                db.session.commit()
                flash(f'{count} relatório(s) importado(s). Arquivos já importados foram ignorados.', 'success')
                return redirect(url_for('patio.listagem'))
            if len(uploads) != 1:
                raise ValueError('Envie os PDFs juntos ou uma planilha por vez.')
            upload = uploads[0]
            rows = ler_relatorio(upload)
            now = datetime.utcnow()
            for data in rows:
                moto = db.session.scalar(db.select(MotoPatio).where(MotoPatio.chassi == data['chassi']))
                if moto is None:
                    moto = MotoPatio(chassi=data['chassi'])
                    db.session.add(moto)
                for key, value in data.items():
                    setattr(moto, key, value)
                moto.atualizado_em = now
                moto.arquivo = Path(upload.filename.replace('\\', '/')).name[:255]
            from inventario_service import sincronizar_fontes
            db.session.flush(); sincronizar_fontes()
            db.session.commit()
            flash(f'Relatório importado: {len(rows)} moto(s) atualizada(s).', 'success')
        except Exception as exc:
            db.session.rollback()
            flash(str(exc) if isinstance(exc, ValueError) else 'Não foi possível importar. Confira o formato do relatório.', 'danger')
        return redirect(url_for('patio.listagem'))
    all_rows, reports = estoque_atual()
    transito = [t for report in reports for t in report.get('transito', [])]
    q = normalizar(request.args.get('q', ''))
    def matches(m):
        grupo = request.args.get('grupo', '')
        grupos = {
            'novas': not m.parcial and normalizar(m.condicao) in ('nova', 'novas', '0 km', '0km'),
            'venda': normalizar(m.situacao) in ('disponivel para venda', 'disponiveis para venda'),
            'aprovacao': normalizar(m.situacao) == 'aguardando aprovacao',
            'atrasadas': atrasada(m),
            'motoshow': 'motoshow' in normalizar(m.faturado_para),
        }
        return ((not q or q in normalizar(' '.join([m.modelo, m.chassi, m.placa, m.cor])))
                and (not request.args.get('situacao') or m.situacao == request.args['situacao'])
                and (not request.args.get('unidade') or m.unidade == request.args['unidade'])
                and (not request.args.get('localizacao') or m.localizacao == request.args['localizacao'])
                and (not grupo or grupos.get(grupo, False)))
    rows = [m for m in all_rows if matches(m)]
    counts = dict(novas=sum(normalizar(m.condicao) == 'nova' and not m.parcial for m in all_rows), venda=sum(normalizar(m.situacao) == 'disponivel para venda' for m in all_rows), aprovacao=sum(normalizar(m.situacao) == 'aguardando aprovacao' for m in all_rows), motoshow=sum('motoshow' in normalizar(m.faturado_para) for m in all_rows))
    unit = request.args.get('unidade', '')
    physical = [m for m in all_rows if not m.parcial and (not unit or m.unidade == unit)]
    counts.update(dentro=sum(m.localizacao == 'Dentro da oficina' for m in physical), transito=sum(m.localizacao == 'Em trânsito' for m in physical), confirmar=sum(m.localizacao == 'A confirmar' for m in physical))
    transito = [t for t in transito if (not unit or t['unidade'] == unit) and (not q or q in normalizar(t['chassi'] + ' ' + t['titulo']))]
    from inventario_service import indicadores, papel
    from inventario_models import MotoEstoque
    master_rows=db.session.scalars(db.select(MotoEstoque).where(MotoEstoque.unidade_nome==unit) if unit else db.select(MotoEstoque)).all()
    return render_template('patio.html', page='patio', rows=rows, total=len(all_rows), counts=counts, reports=reports, transito=transito, indicadores=indicadores(master_rows), papel_estoque=papel(), situacoes=sorted({m.situacao for m in all_rows}), patio_unidades=sorted({m.unidade for m in all_rows if m.unidade}), ultima=max((m.atualizado_em for m in all_rows), default=None))
