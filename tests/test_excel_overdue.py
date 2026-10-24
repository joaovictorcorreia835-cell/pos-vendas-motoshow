import unittest,io
from datetime import datetime,date,timedelta
from openpyxl import load_workbook
import test_app as fixtures
from models import db,Pedido

class ExcelOverdueTest(unittest.TestCase):
 setUp=fixtures.OficinaTest.setUp
 tearDown=fixtures.OficinaTest.tearDown
 form=fixtures.OficinaTest.form
 post=fixtures.OficinaTest.post
 def test_orders_export_uses_filtered_overdue_with_numeric_money(self):
  old=(date.today()-timedelta(days=31)).isoformat();recent=(date.today()-timedelta(days=20)).isoformat()
  self.post(os='0001',data=old,cliente='=1+1');self.post(os='0002',data=recent,cliente='=1+1');self.post(os='0003',data=old,status='Finalizada',cliente='=1+1')
  query='/ordens?inicio=1900-01-01&fim='+date.today().isoformat()+'&prazo=atrasado'
  self.assertIn('Baixar Excel',self.c.get(query).text)
  response=self.c.get(query+'&export=xlsx');self.assertEqual(response.status_code,200);self.assertIn('os_atrasadas',response.headers['Content-Disposition'])
  ws=load_workbook(io.BytesIO(response.data)).active
  self.assertEqual(ws.max_row,3);self.assertEqual(ws['A2'].value,'0001');self.assertEqual(ws['C2'].value,'=1+1');self.assertEqual(ws['C2'].data_type,'s')
  self.assertIsInstance(ws['B2'].value,datetime);self.assertEqual(ws['L2'].value,31);self.assertEqual(ws['O2'].value,160.3);self.assertEqual(ws['O3'].value,'=SUM(O2:O2)')
  self.assertEqual(ws.freeze_panes,'A2');self.assertEqual(ws.auto_filter.ref,'A1:P2')
 def test_pedidos_export_and_filter_retention(self):
  with self.app.app_context():
   for n,days,status,unit in [('1',31,'Aberto','Parauapebas'),('2',20,'Aberto','Parauapebas'),('3',35,'Concluído','Parauapebas'),('4',32,'Aberto','Outra')]:
    db.session.add(Pedido(numero=n,unidade=unit,cliente='Cliente',consultor='Ana',inserido_em=datetime.combine(date.today()-timedelta(days=days),datetime.min.time()),status=status,bruto_centavos=12345,liquido_centavos=10000,codigo_interno=n,arquivo='teste',sha256='a'*64,dados={'Telefone':'001234'}))
   db.session.commit()
  query='/pedidos?prazo=atrasado&unidade=Parauapebas'
  html=self.c.get(query).text;self.assertIn('name="prazo" value="atrasado"',html);self.assertIn('export=xlsx',html)
  response=self.c.get(query+'&export=xlsx');ws=load_workbook(io.BytesIO(response.data)).active
  self.assertEqual(ws.max_row,3);self.assertEqual(ws['A2'].value,'1');self.assertEqual(ws['I2'].value,31);self.assertEqual(ws['L2'].value,100);self.assertEqual(ws['E2'].value,'001234')
  self.assertEqual(ws['L3'].value,'=SUM(L2:L2)')
 def test_empty_export_valid(self):
  r=self.c.get('/pedidos?prazo=atrasado&export=xlsx');ws=load_workbook(io.BytesIO(r.data)).active;self.assertEqual(ws.max_row,1)

if __name__=='__main__':unittest.main()
