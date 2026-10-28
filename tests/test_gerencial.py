import io
import json
import unittest
from copy import deepcopy
from datetime import date, datetime
from unittest.mock import patch
from sqlalchemy import event
from openpyxl import load_workbook
import test_app as fixtures
from models import db, Cliente, Consultor, Passagem, OrdemServico, OrigemOS, ItemOrdem, VendaBalcao, Usuario, FaturamentoMensal, FaturamentoConsultor
from gerencial_models import registrar_balcao, DevolucaoGerencial, FonteGerencial
from gerencial_service import carregar, intervalo, modelo_base


class GerencialTest(unittest.TestCase):
    setUp=fixtures.OficinaTest.setUp
    tearDown=fixtures.OficinaTest.tearDown

    def os(self, numero, entrada='2026-10-05', status='Finalizada', modelo='FZ25 FAZER AZUL 26/26', pecas=10000, servicos=5000, competencia='2026-10', tipo='CLI - SERVIÇOS', cliente=None):
        consultor=db.session.scalar(db.select(Consultor).where(Consultor.nome=='ANA'))
        if not consultor:
            consultor=Consultor(nome='ANA');db.session.add(consultor)
        p=Passagem(cliente=cliente or Cliente(nome='Cliente OS'),data=datetime.fromisoformat(entrada),
            tipo='Manutenção',status=status,unidade='Parauapebas',consultor=consultor,modelo=modelo)
        p.itens=[ItemOrdem(categoria='Peças (Balcão)',centavos=pecas),ItemOrdem(categoria='Mão de obra',centavos=servicos)]
        p.ordem=OrdemServico(numero=str(numero),origem=OrigemOS(arquivo='os.pdf',sha256='a'*64,pagina=1,
            dados=dict(tipo_original=tipo,competencia_faturamento=competencia)))
        db.session.add(p);db.session.flush()
        return p

    def venda(self, nota='NF-e - 1', devolucao=False, custo=True):
        row=dict(numero=nota[-1],codigo_interno='10',cliente='Cliente balcão',consultor='Ana',
            natureza='Devolução' if devolucao else 'Venda',emissao='2026-10-10',data='2026-10-10',nota=nota,
            bruto=-2500 if devolucao else 25000,desconto=-500 if devolucao else 5000,
            liquido=-2000 if devolucao else 20000,pagina=1)
        if custo:
            row['custo']=-1200 if devolucao else 12000
        p=Passagem(cliente=Cliente(nome='Cliente balcão'),data=datetime(2026,10,10),tipo='Balcão',
            status='Finalizada',unidade='Parauapebas')
        p.itens=[ItemOrdem(categoria='Peças (Balcão)',centavos=abs(row['liquido']))]
        p.venda_balcao=VendaBalcao(unidade=p.unidade,numero=row['numero'],codigo_interno='10',nota=nota,
            natureza=row['natureza'],arquivo='balcao.pdf',sha256='b'*64,pagina=1,dados=row)
        db.session.add(p);db.session.flush()
        return row,p

    def relatorio(self, args=None):
        with patch('gerencial_service.hoje_local',return_value=date(2026,10,28)):
            return carregar(args or dict(periodo='mes'))

    def test_receita_carteira_ticket_e_devolucao_sem_duplicar(self):
        with self.app.app_context():
            self.os(1)
            self.os(2,entrada='2026-09-20',modelo='FZ25 FAZER BRANCA 25/26',pecas=5000)
            self.os(3,status='Aberta',modelo='R15 ABS PRETA 24/24',pecas=20000,servicos=1000,entrada='2026-10-01')
            venda,_=self.venda()
            devolucao=dict(venda,natureza='Devolução',bruto=-2500,desconto=-500,liquido=-2000,custo=-1200)
            registrar_balcao([venda,devolucao],dict(sha256='b'*64,arquivo='balcao.pdf'), 'Parauapebas')
            db.session.add(FaturamentoMensal(ano=2026,mes=10,unidade='Parauapebas',linhas=[FaturamentoConsultor(codigo='1',nome='MANUAL',os_centavos=99999999)]))
            db.session.commit()
            d=self.relatorio();i=d['indicadores']
            self.assertEqual((i['passagens'],i['os_faturadas'],i['os_abertas']),(2,2,1))
            self.assertEqual((i['faturamento'],i['pecas'],i['servicos'],i['balcao']),(43000,15000,10000,18000))
            self.assertEqual((i['carteira'],i['ticket'],i['margem'],i['margem_pct']),(21000,12500,7200,40))
            self.assertEqual((i['devolucoes'],i['valor_devolvido'],i['impacto_margem']),(1,2000,-800))
            self.assertEqual(d['modelos'][0]['modelo'],'FZ25 FAZER')
            self.assertEqual(sum(g['total'] for g in d['aging']),21000)
            self.assertEqual(sum(r['balcao'] for r in d['consultores']),18000)
            self.assertEqual(len(d['consultores']),1)  # ANA e Ana são a mesma grafia normalizada.

    def test_devolucao_em_duas_fontes_deduplicada_e_importacao_repetida(self):
        with self.app.app_context():
            row,_=self.venda(devolucao=True)
            resumo=dict(arquivo='balcao.pdf',sha256='b'*64)
            registrar_balcao([row],resumo,'Parauapebas')
            db.session.commit()
            self.assertFalse(registrar_balcao([row],resumo,'Parauapebas')['nova_fonte'])
            db.session.commit()
            d=self.relatorio()
            self.assertEqual(d['indicadores']['devolucoes'],1)
            self.assertEqual(db.session.query(DevolucaoGerencial).count(),1)
            self.assertEqual(db.session.query(FonteGerencial).count(),1)

    def test_dias_nao_inventam_data_fiscal_e_filtros_os_nao_atribuem_balcao(self):
        with self.app.app_context():
            self.os(1,entrada='2026-10-28')
            self.venda()
            db.session.commit()
            d=self.relatorio(dict(periodo='hoje'))
            self.assertEqual(d['indicadores']['passagens'],1)
            self.assertIsNone(d['indicadores']['faturamento'])
            self.assertIsNone(d['indicadores']['ticket'])
            d=self.relatorio(dict(periodo='mes',modelo='FZ25 FAZER'))
            self.assertEqual(d['indicadores']['balcao'],0)
            self.assertEqual(d['indicadores']['faturamento'],15000)
            self.assertEqual(self.relatorio(dict(periodo='mes',consultor='outro'))['indicadores']['faturamento'],0)

    def test_homonimos_separados_e_clientes_unidos_por_cpf(self):
        with self.app.app_context():
            self.os(1,cliente=Cliente(nome='Homonimo'))
            self.os(2,cliente=Cliente(nome='Homonimo'))
            db.session.commit()
            d=self.relatorio()
            self.assertEqual(len(d['clientes']),2)
            self.assertNotEqual(d['clientes'][0]['key'],d['clientes'][1]['key'])

    def test_custo_ausente_nao_vira_zero(self):
        with self.app.app_context():
            self.venda(custo=False);db.session.commit()
            d=self.relatorio()
            self.assertIsNone(d['indicadores']['margem'])
            self.assertIsNone(d['indicadores']['margem_pct'])
            self.assertEqual(d['auditoria']['custos_ausentes'],1)

    def test_nf_ja_contabilizada_na_os_nao_duplica_balcao(self):
        with self.app.app_context():
            p=self.os(1)
            dados=dict(p.ordem.origem.dados,nf_pecas='1')
            p.ordem.origem.dados=dados
            self.venda(nota='NF-e - 1');db.session.commit()
            d=self.relatorio()
            self.assertEqual(d['indicadores']['faturamento'],15000)
            self.assertEqual(d['indicadores']['balcao'],0)
            self.assertEqual(len(d['auditoria']['sobreposicoes_fiscais']),1)

    def test_percentuais_abc_e_cortes(self):
        with self.app.app_context():
            self.os(1,modelo='MODELO A',pecas=8000,servicos=0)
            self.os(2,modelo='MODELO B',pecas=1500,servicos=0)
            self.os(3,modelo='MODELO C',pecas=500,servicos=0)
            db.session.commit()
            modelos=self.relatorio()['modelos']
            self.assertEqual([r['classe'] for r in modelos],['A','B','C'])
            self.assertEqual([r['acumulado'] for r in modelos],[80,95,100])

    def test_faixas_aging_abc_e_matriz(self):
        with self.app.app_context():
            for i,day in enumerate([28,26,25,23,22,18,17,8,7],1):
                self.os(i,entrada=f'2026-10-{day:02d}',status='Aberta',tipo='GAR - GARANTIA')
            db.session.commit();d=self.relatorio()
            self.assertEqual([r['quantidade'] for r in d['aging']],[2,2,2,2,1])
            self.assertEqual(d['indicadores']['garantias'],9)
            self.assertEqual(len(d['matriz'][0]['celulas']['Garantia']),9)

    def test_comparacao_exige_historico_e_usar_competencia(self):
        with self.app.app_context():
            self.os(1,pecas=15000)
            db.session.commit()
            self.assertIsNone(self.relatorio()['comparacoes']['faturamento'])
            self.os(2,entrada='2026-09-02',competencia='2026-09',pecas=5000)
            db.session.commit()
            d=self.relatorio()
            self.assertEqual(d['comparacoes']['faturamento'],100)

    def test_tela_exportacoes_e_sem_gravacao_operacional(self):
        with self.app.app_context():
            self.os(1);self.venda();db.session.commit()
            antes={t.name:[tuple(r) for r in db.session.execute(db.select(t).order_by(*t.primary_key.columns))] for t in db.metadata.sorted_tables}
        response=self.c.get('/painel-gerencial')
        self.assertEqual(response.status_code,200)
        self.assertIn('Visão executiva',response.text)
        self.assertLess(response.text.index('href="/passagens"'),response.text.index('href="/painel-gerencial"'))
        self.assertLess(response.text.index('href="/painel-gerencial"'),response.text.index('href="/faturamento"'))
        r=self.c.get('/painel-gerencial?export=xlsx')
        book=load_workbook(io.BytesIO(r.data))
        self.assertIn('Dados e filtros',book.sheetnames)
        self.assertIn('modelos',book.sheetnames)
        r=self.c.get('/painel-gerencial?export=pdf')
        self.assertEqual(r.status_code,200)
        self.assertTrue(r.data.startswith(b'%PDF'))
        with self.app.app_context():
            for t in db.metadata.sorted_tables:
                self.assertEqual(antes[t.name],[tuple(r) for r in db.session.execute(db.select(t).order_by(*t.primary_key.columns))],t.name)

    def test_permissoes_e_periodo_invalido(self):
        self.assertEqual(self.c.get('/painel-gerencial?periodo=personalizado&inicio=bad&fim=bad').status_code,400)
        with self.app.app_context():
            from inventario_models import UnidadeEstoque
            unidade=db.session.scalar(db.select(UnidadeEstoque))
            usuario=Usuario(login='estoque',senha_hash='unused',perfil='ESTOQUE_MOTOS',unidade_id=unidade.id)
            db.session.add(usuario);db.session.commit();ident=usuario.id
        with self.c.session_transaction() as s:s.update(usuario_id=ident,usuario_versao=1)
        self.assertEqual(self.c.get('/painel-gerencial').status_code,403)
        self.assertEqual(self.c.get('/api/painel-gerencial').status_code,403)

    def test_consultas_nao_crescem_com_cada_os(self):
        with self.app.app_context():
            for i in range(50):self.os(i)
            db.session.commit();db.session.remove()
            consultas=[]
            def contar(*args):consultas.append(1)
            event.listen(db.engine,'before_cursor_execute',contar)
            try:self.relatorio()
            finally:event.remove(db.engine,'before_cursor_execute',contar)
            self.assertLess(len(consultas),20)


class GerencialRegrasTest(unittest.TestCase):
    def test_modelo_nao_depende_de_lista_fixa(self):
        self.assertEqual(modelo_base('MODELO NOVO AZUL METALICO 26/27'),'MODELO NOVO')
        self.assertEqual(modelo_base('YZF R3 ABS'),'YZF R3 ABS')
    def test_intervalos_mes_e_semana(self):
        a,b,c,d=intervalo(dict(periodo='mes'),date(2026,10,28))
        self.assertEqual((a,b,c,d),(date(2026,10,1),date(2026,10,31),date(2026,9,1),date(2026,9,30)))
        self.assertEqual(intervalo(dict(periodo='semana'),date(2026,10,28))[:2],(date(2026,10,26),date(2026,10,28)))
