import unittest, tempfile, io, sys
from pathlib import Path
from datetime import date, timedelta
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import app as module
from app import create_app
from models import db, Passagem, Cliente, Usuario
from werkzeug.security import generate_password_hash
from openpyxl import load_workbook

class OficinaTest(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory(); self.oldroot=module.ROOT; module.ROOT=Path(self.tmp.name)
  self.app=create_app('sqlite:///'+str(Path(self.tmp.name)/'test.db')); self.app.config['TESTING']=True
  with self.app.app_context():
   user=Usuario(login='teste',senha_hash=generate_password_hash('senha-de-teste-123')); db.session.add(user); db.session.commit(); uid=user.id
  self.c=self.app.test_client()
  with self.c.session_transaction() as s: s.update(usuario_id=uid,usuario_versao=1)
  self.c.get('/')
  with self.c.session_transaction() as s: self.token=s['csrf']
 def tearDown(self):
  with self.app.app_context(): db.session.remove(); db.engine.dispose()
  module.ROOT=self.oldroot; self.tmp.cleanup()
 def form(self,**kw):
  d=dict(csrf=self.token,data=date.today().isoformat(),hora='10:30',cliente='Cliente de teste',telefone='94999999999',cpf='12345678901',tipo='Revisão',status='Aberta',unidade='Parauapebas',os='OS-001',modelo='Fazer',placa='ABC1D23',ano='2025',km='100',consultor='Ana',mecanico='José',valor0='100.10',valor1='50.20',valor2='10',valor3='0',valor4='0',valor5='0')
  d.update(kw); return d
 def post(self,**kw): return self.c.post('/lancamento',data=self.form(**kw))
 def data(self,query=''): return self.c.get('/api/dashboard'+query).get_json()
 def test_pages_empty_offline_assets(self):
  for url in ['/','/lancamento','/ordens','/passagens','/faturamento','/pecas','/garantia','/seguro','/clientes','/relatorios','/configuracoes','/static/vendor/bootstrap.min.css','/static/vendor/chart.umd.min.js']:
   response=self.c.get(url); self.assertEqual(response.status_code,200,url); response.close()
  self.assertEqual(self.data()['totals'],[0]*8)
 def test_create_metrics_edit_delete(self):
  self.assertEqual(self.post().status_code,302)
  self.assertEqual(self.data()['totals'],[1,16030,16030,0,5020,1,0,0])
  self.assertAlmostEqual(sum(s[-1] for s in self.data()['series']),160.30)
  self.assertEqual(self.c.get('/lancamento/1').status_code,200)
  self.assertEqual(self.c.post('/lancamento/1',data=self.form(valor0='200',status='Finalizada')).status_code,302)
  self.assertEqual(self.data()['totals'][5],0)
  self.assertEqual(self.data()['totals'][1],26020)
  self.assertEqual(self.c.post('/excluir/1',data={'csrf':self.token}).status_code,302)
  self.assertEqual(self.data()['totals'],[0]*8)
 def test_no_os_and_counter(self):
  self.assertEqual(self.post(os='',valor0='0',valor1='0',valor2='0').status_code,302)
  self.assertEqual(self.data()['totals'][:3],[1,0,0])
  self.assertEqual(self.post(tipo='Balcão',os='').status_code,302)
  self.assertEqual(self.data()['totals'][:4],[1,16030,0,16030])
  with self.app.app_context(): self.assertEqual(db.session.query(Cliente).count(),1)
 def test_validation(self):
  self.post()
  for changes in [dict(os='OS-001'),dict(os='002',valor0='-1'),dict(os='003',cliente=''),dict(os='',valor0='2'),dict(os='005',valor0='NaN'),dict(os='006',km='-1'),dict(os='007',cpf='123')]:
   response=self.post(**changes); self.assertEqual(response.status_code,200)
   with self.app.app_context(): self.assertEqual(db.session.query(Passagem).count(),1)
  self.assertEqual(self.c.post('/lancamento',data={'cliente':'x'}).status_code,400)
  self.assertEqual(self.app.test_client().post('/lancamento',data={'cliente':'x'}).status_code,400)
 def test_filters_previous_and_charts(self):
  self.post(); yesterday=(date.today()-timedelta(days=1)).isoformat(); self.post(os='OS-002',data=yesterday,consultor='Bia')
  self.assertEqual(self.data()['previous'][0],1)
  self.assertEqual(self.data('?periodo=ontem')['totals'][0],1)
  self.assertEqual(self.data('?tipo=Seguro')['totals'],[0]*8)
  self.assertEqual(self.data('?status=Finalizada')['totals'],[0]*8)
  self.assertEqual(self.data('?unidade=Outra')['totals'],[0]*8)
  self.assertEqual(self.data('?consultor=1')['totals'][0],1)
  d=self.data('?inicio='+yesterday+'&fim='+date.today().isoformat()); self.assertEqual(d['totals'][0],2); self.assertEqual(len(d['days']),2)
  self.assertEqual(self.c.get('/api/dashboard?inicio=2026-12-31&fim=2026-01-01').status_code,400)
 def test_special_clients_reports_exports(self):
  self.post(tipo='Garantia',garantia_tipo='Fábrica',especial_status='Aprovada',descricao='Reparo')
  self.post(os='OS-002',tipo='Seguro',seguradora='Seguradora',sinistro='123',especial_status='Em análise')
  d=self.data(); self.assertEqual(d['totals'][6:],[1,1]); self.assertEqual(d['types'][2:4],[1,1])
  for kind in ['diario','mensal','passagens','consultor','mecanico','abertas','garantias','seguros','pecas','mao']:
   self.assertEqual(self.c.get('/relatorios?relatorio='+kind).status_code,200)
  xlsx=self.c.get('/relatorios?export=xlsx'); self.assertEqual(xlsx.status_code,200); self.assertEqual(load_workbook(io.BytesIO(xlsx.data)).active.max_row,2)
  pdf=self.c.get('/relatorios?export=pdf'); self.assertEqual(pdf.status_code,200); self.assertTrue(pdf.data.startswith(b'%PDF'))
  self.assertEqual(self.c.get('/clientes?id=1').status_code,200)
  self.c.post('/clientes/1/excluir',data={'csrf':self.token})
  with self.app.app_context(): self.assertEqual(db.session.query(Cliente).count(),1)
  self.c.post('/clientes',data=dict(csrf=self.token,nome='Outro',cpf='98765432100',telefone='111'))
  self.c.post('/clientes',data=dict(csrf=self.token,id='2',nome='Outro editado',cpf='98765432100'))
  self.c.post('/clientes/2/excluir',data={'csrf':self.token})
  with self.app.app_context(): self.assertEqual(db.session.query(Cliente).count(),1)
 def test_backup_restore(self):
  self.post(); backup=self.c.post('/backup',data={'csrf':self.token}); self.assertEqual(backup.status_code,200); original=backup.data; backup.close()
  self.c.post('/excluir/1',data={'csrf':self.token}); self.assertEqual(self.data()['totals'][0],0)
  response=self.c.post('/restaurar',data={'csrf':self.token,'confirmar':'SIM','arquivo':(io.BytesIO(original),'backup.db')},content_type='multipart/form-data')
  self.assertEqual(response.status_code,302); self.assertEqual(self.data()['totals'][0],1)
  self.c.post('/restaurar',data={'csrf':self.token,'confirmar':'SIM','arquivo':(io.BytesIO(b'bad'),'bad.db')},content_type='multipart/form-data')
  self.assertEqual(self.data()['totals'][0],1)
 def test_cancelled_not_revenue(self):
  self.post(status='Cancelada'); self.assertEqual(self.data()['totals'],[0]*8); self.assertEqual(sum(self.data()['categories']),0)

if __name__=='__main__': unittest.main(verbosity=2)
